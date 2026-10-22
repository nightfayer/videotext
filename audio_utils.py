import asyncio
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

def sanitize_filename(name: str, max_length: int = 100) -> str:
    """Removes invalid OS/filesystem characters and limits filename length."""
    if not name:
        return f"file_{int(time.time())}"
    # Replace invalid Windows/Linux filename characters: \ / : * ? " < > |
    clean = re.sub(r'[\\/*?:"<>|\x00-\x1f]', '_', str(name)).strip(" ._")
    if not clean:
        return "media"
    return clean[:max_length]


def format_duration(seconds: float) -> str:
    """Formats seconds into human-readable MM:SS or HH:MM:SS."""
    total_seconds = int(seconds)
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    secs = total_seconds % 60
    if hours > 0:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def format_bytes(bytes_count: int) -> str:
    """Formats bytes into KB, MB, GB."""
    if bytes_count < 1024:
        return f"{bytes_count} B"
    elif bytes_count < 1024 * 1024:
        return f"{bytes_count / 1024:.1f} KB"
    elif bytes_count < 1024 * 1024 * 1024:
        return f"{bytes_count / (1024 * 1024):.1f} МБ"
    else:
        return f"{bytes_count / (1024 * 1024 * 1024):.2f} ГБ"


_FFMPEG_PATH: str | None = None


def is_ffmpeg_functional(ffmpeg_path: str) -> bool:
    """Checks if the given ffmpeg executable actually runs."""
    try:
        res = subprocess.run(
            [ffmpeg_path, "-version"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=5
        )
        return res.returncode == 0
    except Exception:
        return False


def install_ffmpeg() -> str | None:
    """
    Attempts to automatically install FFmpeg:
    1. Installs imageio-ffmpeg via pip (cross-platform, rootless, works on Windows/Linux/macOS).
    2. On Linux, tries apt-get update && apt-get install -y ffmpeg.
    3. On Windows, tries winget install Gyan.FFmpeg.
    """
    print("[FFmpeg Auto-Install] FFmpeg не найден на устройстве. Запуск автоматической установки...", flush=True)

    # 1. Cross-platform via pip (imageio-ffmpeg bundles standalone ffmpeg)
    try:
        print("[FFmpeg Auto-Install] 1/3: Установка пакета imageio-ffmpeg через pip...", flush=True)
        subprocess.run(
            [sys.executable, "-m", "pip", "install", "imageio-ffmpeg"],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE
        )
        import imageio_ffmpeg
        exe = imageio_ffmpeg.get_ffmpeg_exe()
        if exe and is_ffmpeg_functional(exe):
            print(f"[FFmpeg Auto-Install] ✅ Успешно установлен (imageio-ffmpeg): {exe}", flush=True)
            return exe
    except Exception as e:
        print(f"[FFmpeg Auto-Install] Установка через pip не удалась ({e}). Пробуем системные менеджеры пакетов...", flush=True)

    # 2. Linux apt-get
    if sys.platform.startswith("linux"):
        try:
            print("[FFmpeg Auto-Install] 2/3: Установка через apt-get...", flush=True)
            is_root = (os.geteuid() == 0) if hasattr(os, "geteuid") else False
            prefix = [] if is_root else ["sudo"]
            subprocess.run(prefix + ["apt-get", "update"], check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            subprocess.run(prefix + ["apt-get", "install", "-y", "ffmpeg"], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            found = shutil.which("ffmpeg")
            if found and is_ffmpeg_functional(found):
                print(f"[FFmpeg Auto-Install] ✅ Успешно установлен через apt: {found}", flush=True)
                return found
        except Exception as e:
            print(f"[FFmpeg Auto-Install] Ошибка apt-get: {e}", flush=True)

    # 3. Windows winget
    if sys.platform == "win32":
        try:
            print("[FFmpeg Auto-Install] 3/3: Установка Gyan.FFmpeg через winget...", flush=True)
            subprocess.run(
                ["winget", "install", "Gyan.FFmpeg", "--accept-package-agreements", "--accept-source-agreements", "--silent"],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )
            appdata_local = os.environ.get("LOCALAPPDATA", "")
            if appdata_local:
                winget_path = Path(appdata_local) / "Microsoft" / "WinGet" / "Packages"
                if winget_path.exists():
                    for match in winget_path.glob("**/ffmpeg.exe"):
                        if is_ffmpeg_functional(str(match)):
                            print(f"[FFmpeg Auto-Install] ✅ Успешно найден после winget: {match}", flush=True)
                            return str(match)
        except Exception as e:
            print(f"[FFmpeg Auto-Install] Ошибка winget: {e}", flush=True)

    return None


def get_ffmpeg_executable(auto_install: bool = True) -> str:
    """
    Finds the ffmpeg executable in PATH, imageio_ffmpeg, or common locations.
    If not found and auto_install=True, automatically installs it.
    """
    global _FFMPEG_PATH
    if _FFMPEG_PATH and is_ffmpeg_functional(_FFMPEG_PATH):
        return _FFMPEG_PATH

    # 1. System PATH
    found = shutil.which("ffmpeg")
    if found and is_ffmpeg_functional(found):
        _FFMPEG_PATH = found
        return _FFMPEG_PATH

    # 2. imageio_ffmpeg (pip package with bundled ffmpeg binary)
    try:
        import imageio_ffmpeg
        exe = imageio_ffmpeg.get_ffmpeg_exe()
        if exe and is_ffmpeg_functional(exe):
            _FFMPEG_PATH = exe
            return _FFMPEG_PATH
    except (ImportError, Exception):
        pass

    # 3. Common Windows locations
    if sys.platform == "win32":
        appdata_local = os.environ.get("LOCALAPPDATA", "")
        if appdata_local:
            winget_path = Path(appdata_local) / "Microsoft" / "WinGet" / "Packages"
            if winget_path.exists():
                for match in winget_path.glob("**/ffmpeg.exe"):
                    if is_ffmpeg_functional(str(match)):
                        _FFMPEG_PATH = str(match)
                        return _FFMPEG_PATH

        candidates = [
            Path("ffmpeg.exe"),
            Path("bin/ffmpeg.exe"),
            Path(r"C:\ffmpeg\bin\ffmpeg.exe"),
            Path(r"C:\Program Files\ffmpeg\bin\ffmpeg.exe"),
        ]
        for c in candidates:
            if c.exists() and is_ffmpeg_functional(str(c)):
                _FFMPEG_PATH = str(c)
                return _FFMPEG_PATH

    # 4. If still not found, run automatic installation
    if auto_install:
        installed = install_ffmpeg()
        if installed:
            _FFMPEG_PATH = installed
            return _FFMPEG_PATH

    return "ffmpeg"


def extract_audio_ffmpeg(input_file: Path, output_file: Path) -> tuple[bool, str]:
    """
    Extracts audio from video/audio file into 16kHz mono WAV using ffmpeg.
    Returns (success, error_message).
    """
    if not input_file.exists():
        return False, "Скачанный файл не найден на сервере."

    file_size = input_file.stat().st_size
    if file_size == 0:
        return False, "Скачанный файл пуст (0 байт). Возможно, сбой при загрузке."

    ffmpeg_bin = get_ffmpeg_executable()
    cmd = [
        ffmpeg_bin,
        "-y",
        "-threads", "4",
        "-i", str(input_file),
        "-vn",
        "-c:a", "pcm_s16le",
        "-ac", "1",
        "-ar", "16000",
        "-nostats",
        "-loglevel", "error",
        "-f", "wav",
        str(output_file)
    ]
    try:
        subprocess.run(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=True
        )
        if not output_file.exists() or output_file.stat().st_size == 0:
            return False, "Файл звуковой дорожки не создался или пуст."
        return True, ""
    except FileNotFoundError:
        if os.name == "nt":
            msg = (
                "Утилита FFmpeg не найдена в системе!\n"
                "Если вы только что установили FFmpeg, перезапустите бота (закройте окно и запустите снова run.bat).\n"
                "Если не установлена: выполните в PowerShell: winget install Gyan.FFmpeg"
            )
        else:
            msg = "Утилита FFmpeg не найдена на сервере! Установите её: apt update && apt install -y ffmpeg"
        print(f"[FFmpeg Missing]: {msg}", flush=True)
        return False, msg
    except subprocess.CalledProcessError as e:
        err = e.stderr.decode("utf-8", errors="ignore").strip()
        print(f"[FFmpeg Error]: {err}", flush=True)
        if "does not contain any stream" in err or "Output file is empty" in err:
            return False, "В этом видео отсутствует звуковая дорожка (нет звука)."
        return False, f"Ошибка FFmpeg при обработке файла: {err[:150]}"
    except Exception as e:
        err = str(e)
        print(f"[Audio Extraction Error]: {err}", flush=True)
        return False, f"Ошибка извлечения звука: {err}"


async def async_extract_audio(input_file: Path, output_file: Path) -> tuple[bool, str]:
    """Runs audio extraction in an async thread pool. Returns (success, error_message)."""
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(None, extract_audio_ffmpeg, input_file, output_file)


def safe_remove(file_path: Path):
    """Safely removes a file if it exists."""
    try:
        if file_path and file_path.exists():
            file_path.unlink()
    except Exception as e:
        print(f"[Cleanup Warning]: Could not remove {file_path}: {e}")


def clean_temp_dir(temp_dir: Path):
    """Cleans up leftover temporary files on bot startup."""
    try:
        if not temp_dir.exists():
            temp_dir.mkdir(parents=True, exist_ok=True)
            return
        cleaned = 0
        for item in temp_dir.iterdir():
            if item.is_file():
                try:
                    item.unlink()
                    cleaned += 1
                except Exception:
                    pass
        if cleaned > 0:
            print(f"🧹 Очищено старых временных файлов: {cleaned}", flush=True)
    except Exception as e:
        print(f"[Temp Cleanup Error]: {e}", flush=True)


def create_docx_report(
    base_name: str,
    duration_str: str,
    lang_name: str,
    prob: int,
    proc_time: float,
    text: str,
    timestamped_text: str,
    output_path: Path
) -> bool:
    """Generates a styled Microsoft Word (.docx) document with transcription."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()

        # Document Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(f"Транскрибация: {base_name}")
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)
        title_p.paragraph_format.space_after = Pt(12)

        # Metadata block
        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_after = Pt(16)
        r1 = meta_p.add_run(f"⏱️ Длительность: {duration_str}\n")
        r1.font.bold = True
        r2 = meta_p.add_run(f"🌐 Язык: {lang_name} ({prob}%)\n")
        r3 = meta_p.add_run(f"⚡ Время обработки: {proc_time:.1f} сек\n")

        # Headings & Text
        h1 = doc.add_heading("Текст с таймкодами", level=2)
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(6)
        p_ts = doc.add_paragraph(timestamped_text)
        p_ts.paragraph_format.space_after = Pt(16)

        h2 = doc.add_heading("Полный текст", level=2)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(6)
        doc.add_paragraph(text)

        doc.save(str(output_path))
        return True
    except Exception as e:
        print(f"[DOCX Generation Error]: {e}", flush=True)
        return False
